"""cv templates primitives."""

from __future__ import annotations

import hashlib
import io
import re

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

from . import docx_safety

PROFILE_HEADINGS = {
    "PROFIEL",
    "PROFIEL EN AMBITIE",
    "OVER MIJ",
    "SAMENVATTING",
    "PROFILE",
    "SUMMARY",
}


STRONG_HEADINGS = {"STERKE PUNTEN", "KERNKWALITEITEN", "COMPETENTIES", "VAARDIGHEDEN", "SKILLS"}


def _all_paragraphs(document):
    """Yield every paragraph, including table cells (the CV layout is tabular)."""
    for element in document.element.body.iter(qn("w:p")):
        yield Paragraph(element, document)


def _set_text(paragraph, text: str) -> None:
    """Replace paragraph text while keeping the first run's formatting."""
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def _is_list(paragraph) -> bool:
    properties = paragraph._p.pPr
    if properties is not None and properties.numPr is not None:
        num_id = properties.numPr.numId
        return num_id is None or num_id.val != 0
    style = paragraph.style
    seen = set()
    while style is not None:
        if style.style_id in seen:
            return False
        seen.add(style.style_id)
        if style.name.startswith(("List Bullet", "List Number")):
            return True
        properties = style.element.pPr
        if properties is not None and properties.numPr is not None:
            num_id = properties.numPr.numId
            return num_id is None or num_id.val != 0
        style = style.base_style
    return False


def inspect_cv(raw: bytes) -> dict:
    """Inspect without writes; return structure only, never personal paragraph text."""
    docx_safety.validate_package(raw)
    try:
        document = Document(io.BytesIO(raw))
    except Exception:
        raise ValueError("Geen leesbaar DOCX-bestand. Kies een geldig bestand opnieuw.") from None
    paragraphs = list(_all_paragraphs(document))
    profiles = [i for i, p in enumerate(paragraphs) if p.text.strip().upper() in PROFILE_HEADINGS]
    strengths = [i for i, p in enumerate(paragraphs) if p.text.strip().upper() in STRONG_HEADINGS]
    errors, profile_index, strong_indices = [], None, []
    if len(profiles) != 1:
        errors.append("Gebruik één herkenbare profielkop, bijvoorbeeld 'Profiel' of 'Over mij'.")
    else:
        for index in range(profiles[0] + 1, len(paragraphs)):
            paragraph = paragraphs[index]
            if not paragraph.text.strip():
                continue
            if index in strengths or paragraph.style.name.startswith("Heading"):
                break
            profile_index = index
            break
        if profile_index is None:
            errors.append("De profielkop is gevonden, maar profieltekst ontbreekt.")
    if len(strengths) != 1:
        errors.append("Gebruik één kop 'Sterke punten' of 'Kernkwaliteiten', met 3–8 lijstpunten.")
    else:
        for index in range(strengths[0] + 1, len(paragraphs)):
            paragraph = paragraphs[index]
            if not paragraph.text.strip():
                continue
            if not _is_list(paragraph) or paragraph.style.name.startswith("Heading"):
                break
            strong_indices.append(index)
        if not 3 <= len(strong_indices) <= 8:
            errors.append(
                f"{len(strong_indices)} sterke punten herkend; toegestaan zijn 3–8 lijstpunten."
            )
    if errors:
        raise ValueError(" ".join(errors))
    return {
        "profile_heading": paragraphs[profiles[0]].text.strip(),
        "strong_heading": paragraphs[strengths[0]].text.strip(),
        "strong_count": len(strong_indices),
        "profile_index": profile_index,
        "strong_indices": strong_indices,
        "source_hash": hashlib.sha256(raw).hexdigest(),
    }


def template_count(document) -> int:
    texts = [p.text.strip() for p in _all_paragraphs(document)]
    indices = [
        int(match.group(1))
        for text in texts
        if (match := re.fullmatch(r"\{\{STRONG_(\d+)\}\}", text))
    ]
    if (
        not 3 <= len(indices) <= 8
        or sorted(indices) != list(range(1, len(indices) + 1))
        or texts.count("{{PROFILE}}") != 1
        or sum("{{STRONG_" in text for text in texts) != len(indices)
    ):
        raise ValueError(
            "CV template placeholders must contain one profile and 3–8 unique contiguous strong points"
        )
    return len(indices)


def prepare_template(raw: bytes) -> tuple[bytes, dict]:
    inspection = inspect_cv(raw)
    document = Document(io.BytesIO(raw))
    identity = _extract_identity(document)
    paragraphs = list(_all_paragraphs(document))
    _set_text(paragraphs[inspection["profile_index"]], "{{PROFILE}}")
    for number, index in enumerate(inspection["strong_indices"], 1):
        _set_text(paragraphs[index], f"{{{{STRONG_{number}}}}}")
    _apply_a4(document)
    template_count(document)
    output = io.BytesIO()
    document.save(output)
    prepared = output.getvalue()
    return prepared, {
        **identity,
        "strong_count": inspection["strong_count"],
        "template_hash": hashlib.sha256(prepared).hexdigest(),
    }


def _extract_identity(document) -> dict:
    identity: dict[str, list[str]] = {"name": [], "tagline": [], "contact": []}
    for paragraph in _all_paragraphs(document):
        text = paragraph.text.strip()
        if text.upper() in PROFILE_HEADINGS:
            break
        if not text:
            continue
        if not identity["name"]:
            identity["name"].append(text)
        elif not identity["tagline"]:
            identity["tagline"].append(text)
        else:
            identity["contact"].append(text)
    return identity


def _is_a4(document) -> bool:
    section = document.sections[0]
    return abs(section.page_width - Cm(21.0)) < 1000 and abs(section.page_height - Cm(29.7)) < 1000


def _apply_a4(document) -> None:
    for section in document.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
