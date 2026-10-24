"""Tailored CV + motivation letter generation, with docx -> pdf conversion."""

from __future__ import annotations

import hashlib
import io
import json
import math
import re
import subprocess
import tempfile
import threading
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from . import config, docx_safety, llm, llm_schemas, profile, standards, store, tenancy
from .cv_templates import (  # noqa: F401
    PROFILE_HEADINGS,
    STRONG_HEADINGS,
    _all_paragraphs,
    _apply_a4,
    _extract_identity,
    _is_a4,
    _is_list,
    _set_text,
    inspect_cv,
    prepare_template,
    template_count,
)
from .document_conversion import _soffice, convert_pdf  # noqa: F401
from .util import now_iso, now_local, parse_json_object

DOC_TEMPLATES_DIR = config.DATA_DIR / "templates"
CV_TEMPLATE = DOC_TEMPLATES_DIR / "cv-template.docx"
IDENTITY_FILE = DOC_TEMPLATES_DIR / "identity.json"

LETTER_MAX_LINES = 15
LETTER_CHARS_PER_LINE = 95
_GENERATION_LOCKS: dict[tuple[str, int, str], threading.Lock] = {}
_GENERATION_LOCKS_GUARD = threading.Lock()
_TEMPLATE_LOCKS: dict[str, threading.RLock] = {}


def template_lock():
    with _GENERATION_LOCKS_GUARD:
        return _TEMPLATE_LOCKS.setdefault(tenancy.user_id(), threading.RLock())


def _generation_lock(owner: str, job_id: int, kind: str) -> threading.Lock:
    key = (owner, job_id, kind)
    with _GENERATION_LOCKS_GUARD:
        return _GENERATION_LOCKS.setdefault(key, threading.Lock())


# ---------------------------------------------------------------- templates


def _atomic_write(path: Path, raw: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=path.parent, delete=False) as temporary:
        name = Path(temporary.name)
        try:
            temporary.write(raw)
            temporary.close()
            name.replace(path)
        finally:
            name.unlink(missing_ok=True)


def ensure_templates() -> None:
    with template_lock():
        target = config.user_path(CV_TEMPLATE, "templates")
        if target.exists():
            raw = target.read_bytes()
            docx_safety.validate_package(raw)
            document = Document(io.BytesIO(raw))
            if not _is_a4(document):
                _apply_a4(document)
                output = io.BytesIO()
                document.save(output)
                raw = output.getvalue()
                _atomic_write(target, raw)
            identity = {
                **_extract_identity(document),
                "strong_count": template_count(document),
                "template_hash": hashlib.sha256(raw).hexdigest(),
            }
        else:
            source = config.seed_dir() / config.seed_files()["cv"]
            raw, identity = prepare_template(source.read_bytes())
            _atomic_write(target, raw)
        identity_path = config.user_path(IDENTITY_FILE, "templates")
        encoded = json.dumps(identity, ensure_ascii=False, indent=2).encode()
        if not identity_path.exists() or identity_path.read_bytes() != encoded:
            _atomic_write(identity_path, encoded)


def cv_identity() -> dict:
    with template_lock():
        ensure_templates()
        return json.loads(config.user_path(IDENTITY_FILE, "templates").read_text())


# ---------------------------------------------------------------- parsing


def slugify(value: str, limit: int = 60) -> str:
    normalized = unicodedata.normalize("NFKD", value or "").encode("ascii", "ignore").decode()
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", normalized).strip("-").lower()
    return (slug or "job")[:limit].strip("-")


def parse_cv_tailoring(text: str, strong_count: int = 5) -> dict:
    data = parse_json_object(text)
    if not isinstance(data.get("strong_points"), list):
        raise ValueError("CV strong_points must be a JSON array")
    if len(data["strong_points"]) != strong_count:
        raise ValueError(f"CV tailoring needs exactly {strong_count} strong points")
    llm_schemas.cv_tailoring_model(strong_count).model_validate(data)
    profile = str(data.get("profile") or "").strip()
    strong = [str(item).strip() for item in (data.get("strong_points") or []) if str(item).strip()]
    if not (120 <= len(profile) <= 800):
        raise ValueError("CV profile must be 120-800 characters")
    if len(strong) != strong_count:
        raise ValueError(f"CV tailoring needs exactly {strong_count} strong points")
    if any(len(item) > 70 for item in strong):
        raise ValueError("strong points must be at most 70 characters")
    return (
        llm_schemas.cv_tailoring_model(strong_count)
        .model_validate({"profile": profile, "strong_points": strong})
        .model_dump()
    )


def parse_letter(text: str) -> list[str]:
    data = parse_json_object(text)
    paragraphs = [str(item).strip() for item in (data.get("paragraphs") or []) if str(item).strip()]
    if not paragraphs:
        raise ValueError("letter needs at least one paragraph")
    visual_lines = sum(max(1, math.ceil(len(item) / LETTER_CHARS_PER_LINE)) for item in paragraphs)
    if visual_lines > LETTER_MAX_LINES:
        raise ValueError(
            f"letter must stay within {LETTER_MAX_LINES} rendered lines "
            f"(estimated {visual_lines}); shorten the text"
        )
    return paragraphs


# ---------------------------------------------------------------- prompts


def cv_prompt(job, profile_md: str, strong_count: int = 5) -> str:
    standards_text = standards.standards_block(("cv-brief-standaarden", "vacature-analyse"))
    extra = f"\n\n{standards_text}" if standards_text else ""
    return (
        f"Pas het cv van {tenancy.display_name()} licht aan voor deze vacature. Gebruik alleen feiten uit het profiel.\n"
        "Verzin geen ervaring, opleiding of vaardigheden.\n\n"
        "## Profiel\n"
        f"{profile_md[:7000]}\n\n"
        "## Vacature\n"
        f"{job['title']} bij {job['company'] or 'onbekende organisatie'}\n"
        f"{job['description'][:7000]}"
        f"{extra}\n\n"
        'Geef alleen JSON: {"profile": "<profielparagraaf van 3-4 zinnen, max 600 tekens>", '
        f'"strong_points": ["<exact {strong_count} korte punten, elk max 60 tekens>"]}}.'
    )


def letter_prompt(job, profile_md: str) -> str:
    standards_text = standards.standards_block(("cv-brief-standaarden", "vacature-analyse"))
    extra = f"\n\n{standards_text}" if standards_text else ""
    return (
        f"Schrijf een korte Nederlandse motivatiebrief voor {tenancy.display_name()}. Maximaal 15 "
        "gerenderde regels in totaal (ongeveer 1.200 tekens), 3-5 alinea's, concreet "
        "en zonder clichés. Gebruik alleen feiten uit het profiel; verzin niets. "
        "Verwijs naar de vacature en waarom deze persoon daar past. Geen adressen, geen "
        "datum, geen ondertekening in de tekst.\n\n"
        "## Profiel\n"
        f"{profile_md[:7000]}\n\n"
        "## Vacature\n"
        f"{job['title']} bij {job['company'] or 'onbekende organisatie'}\n"
        f"{job['description'][:7000]}"
        f"{extra}\n\n"
        'Geef alleen JSON: {"paragraphs": ["<alinea>", "..."]}.'
    )


# ---------------------------------------------------------------- rendering


def _job_dir(job) -> Path:
    company = job["company"] or "onbekend"
    title = job["title"] or "vacature"
    stamp = now_local().strftime("%Y-%m-%d")
    name = f"{stamp}-{slugify(company, 24)}-{slugify(title, 40)}-{job['id']}"
    directory = config.jobs_dir() / name
    directory.mkdir(parents=True, exist_ok=True)
    return directory


def _job_json(job, directory: Path) -> None:
    payload = dict(job)
    payload["generated_at"] = now_iso()
    (directory / "job.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2))


def _next_document_version(directory: Path, kind: str, documents: list) -> int:
    """Never reuse a version, even when an interrupted run left orphan files."""
    versions = {doc["version"] for doc in documents if doc["kind"] == kind}
    pattern = re.compile(rf"^{re.escape(kind)}-v(\d+)\.(?:docx|pdf|json)$")
    versions.update(
        int(match.group(1))
        for path in directory.iterdir()
        if path.is_file() and (match := pattern.match(path.name))
    )
    version = max(versions, default=0) + 1
    while any(
        (directory / f"{kind}-v{version}{suffix}").exists() for suffix in (".docx", ".pdf", ".json")
    ):
        version += 1
    return version


def _record_generation_usage(kind: str, usage: dict) -> None:
    store.record_usage(
        None,
        "documents",
        int(usage.get("input_tokens") or 0),
        int(usage.get("output_tokens") or 0),
        float(usage.get("cost_usd") or 0.0),
        agent=("cv_writer" if kind == "resume" else "letter_writer"),
    )


def render_cv(job, tailoring: dict, out_path: Path, *, template_bytes: bytes | None = None) -> None:
    document = Document(
        io.BytesIO(template_bytes)
        if template_bytes is not None
        else str(config.user_path(CV_TEMPLATE, "templates"))
    )
    count = template_count(document)
    if count != len(tailoring["strong_points"]):
        raise ValueError(f"CV template needs exactly {count} strong points")
    values = {"{{PROFILE}}": tailoring["profile"]}
    values.update(
        {f"{{{{STRONG_{i}}}}}": text for i, text in enumerate(tailoring["strong_points"], 1)}
    )
    replaced = 0
    for paragraph in _all_paragraphs(document):
        key = paragraph.text.strip()
        if key in values:
            _set_text(paragraph, values[key])
            replaced += 1
    if replaced != 1 + count:
        raise RuntimeError(f"CV template placeholders missing (replaced {replaced}/{1 + count})")
    document.save(str(out_path))


def render_letter(job, paragraphs: list[str], out_path: Path) -> None:
    identity = cv_identity()
    name = (identity.get("name") or [tenancy.display_name()])[0]
    contact = [line for line in identity.get("contact", []) if "Rijbewijs" not in line]
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15

    heading = document.add_paragraph()
    run = heading.add_run(name)
    run.bold = True
    run.font.size = Pt(20)
    heading.paragraph_format.space_after = Pt(2)

    if contact:
        line = document.add_paragraph("   |   ".join(contact))
        line.runs[0].font.size = Pt(10)

    document.add_paragraph("")
    city = profile.read_preferences().get("home", {}).get("label", "")
    city = "" if city == "Nog niet ingesteld" else city
    date_line = document.add_paragraph(
        f"{city + ', ' if city else ''}{now_local().strftime('%d-%m-%Y')}"
    )
    date_line.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    employer = document.add_paragraph()
    employer.add_run(job["company"] or "").bold = True
    document.add_paragraph("")

    subject = document.add_paragraph()
    subject.add_run(f"Betreft: sollicitatie naar de functie {job['title']}").bold = True
    document.add_paragraph("")

    for text in paragraphs:
        document.add_paragraph(text)

    document.add_paragraph("")
    document.add_paragraph("Met vriendelijke groet,")
    signature = document.add_paragraph()
    signature.add_run(name).bold = True
    document.save(str(out_path))


# ---------------------------------------------------------------- pdf


# ---------------------------------------------------------------- generation


def _generate_job_document(mgr: llm.ServerManager, job, kind: str) -> dict:
    if kind not in ("resume", "motivatie"):
        raise ValueError("kind must be resume or motivatie")
    ensure_templates()
    with template_lock():
        template_path = config.user_path(CV_TEMPLATE, "templates")
        template_bytes = template_path.read_bytes() if template_path.exists() else None
        strong_count = template_count(Document(io.BytesIO(template_bytes))) if template_bytes else 5
    profile_md = profile.read_public_excerpt(12000)
    directory = _job_dir(job)
    _job_json(job, directory)
    version = _next_document_version(directory, kind, store.list_documents(job["id"]))
    stem = f"{kind}-v{version}"
    docx_path = directory / f"{stem}.docx"
    if docx_path.exists():  # defensive: a version was reserved by an unexpected writer
        raise FileExistsError(f"{docx_path.name} already exists")

    if kind == "resume":
        text, usage = mgr.generate(
            cv_prompt(job, profile_md, strong_count),
            agent="cv_writer",
            schema=llm_schemas.cv_tailoring_model(strong_count).model_json_schema(),
        )
    else:
        text, usage = mgr.generate(letter_prompt(job, profile_md), agent="letter_writer")
    # Account for the paid call before parsing/rendering can fail.
    _record_generation_usage(kind, usage)
    if kind == "resume":
        render_cv(
            job, parse_cv_tailoring(text, strong_count), docx_path, template_bytes=template_bytes
        )
    else:
        render_letter(job, parse_letter(text), docx_path)

    pdf_path = None
    pdf_error = None
    try:
        pdf_path = convert_pdf(docx_path)
    except (subprocess.SubprocessError, OSError) as exc:
        pdf_error = f"{type(exc).__name__}"

    meta = {
        "kind": kind,
        "version": version,
        "job_id": job["id"],
        "created_at": now_iso(),
        "usage": usage,
        "pdf_error": pdf_error,
    }
    (directory / f"{stem}.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2))
    store.add_document(
        job["id"],
        kind,
        version,
        str(docx_path.relative_to(config.DATA_DIR)),
        str(pdf_path.relative_to(config.DATA_DIR)) if pdf_path else None,
        json.dumps(meta, ensure_ascii=False),
    )
    return {
        "docx": str(docx_path),
        "pdf": str(pdf_path) if pdf_path else None,
        "version": version,
        "meta": meta,
    }


def generate_for_job(mgr: llm.ServerManager, job, kind: str) -> dict:
    """Serialize one owner/job/kind without blocking other accounts."""
    with _generation_lock(tenancy.user_id(), job["id"], kind):
        return _generate_job_document(mgr, job, kind)
