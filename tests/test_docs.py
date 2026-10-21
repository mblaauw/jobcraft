from __future__ import annotations

import pytest

from app.docs import LETTER_MAX_LINES, parse_cv_tailoring, parse_letter, slugify


def test_slugify_handles_dutch_titles():
    assert slugify("Stichting Deventer Verhaal") == "stichting-deventer-verhaal"
    assert slugify("Muziekgebouw aan ’t IJ") == "muziekgebouw-aan-t-ij"
    assert slugify("") == "job"


def test_parse_cv_tailoring_requires_five_strong_points():
    payload = {
        "profile": "x" * 200,
        "strong_points": ["a", "b", "c", "d", "e"],
    }
    import json

    parsed = parse_cv_tailoring(json.dumps(payload))
    assert parsed["strong_points"] == ["a", "b", "c", "d", "e"]
    with pytest.raises(ValueError):
        parse_cv_tailoring(json.dumps({**payload, "strong_points": ["a"]}))


def test_letter_line_limit():
    import json

    ok = {"paragraphs": ["regel " + str(i) for i in range(LETTER_MAX_LINES)]}
    assert len(parse_letter(json.dumps(ok))) == LETTER_MAX_LINES
    too_long = {"paragraphs": ["regel " + str(i) for i in range(LETTER_MAX_LINES + 1)]}
    with pytest.raises(ValueError):
        parse_letter(json.dumps(too_long))


def test_letter_char_budget():
    import json

    long_paragraph = "x" * 1500
    with pytest.raises(ValueError):
        parse_letter(json.dumps({"paragraphs": [long_paragraph]}))


def test_letter_rejects_empty():
    import json

    with pytest.raises(ValueError):
        parse_letter(json.dumps({"paragraphs": []}))


def test_identity_stops_at_alternate_profile_heading():
    from docx import Document

    from app import docs

    document = Document()
    document.add_paragraph("Test Person")
    document.add_paragraph("Role")
    document.add_paragraph("OVER MIJ")
    document.add_paragraph("Not identity")
    identity = docs._extract_identity(document)
    assert identity["name"] == ["Test Person"]
