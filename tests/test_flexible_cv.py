import io
import json
import threading
from concurrent.futures import ThreadPoolExecutor

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app import config, cv_uploads, docs, llm, llm_schemas, store, tenancy
from tests.factories import add_job

A = {"id": "a", "username": "a"}
B = {"id": "b", "username": "b"}


def cv_bytes(count=5, style="List Bullet", table=False):
    document = Document()
    parent = document.add_table(rows=1, cols=1).cell(0, 0) if table else document
    if style == "Inherited":
        inherited = document.styles.add_style(style, WD_STYLE_TYPE.PARAGRAPH)
        inherited.base_style = document.styles["List Bullet"]
    parent.add_paragraph("Synthetic Person")
    parent.add_paragraph("Synthetic profession")
    parent.add_paragraph("OVER MIJ")
    parent.add_paragraph("Original public profile from this account.")
    parent.add_paragraph("KERNKWALITEITEN")
    for i in range(count):
        paragraph = parent.add_paragraph(
            f"Synthetic skill {i}", style=None if style == "Numbered" else style
        )
        if style == "Numbered":
            properties = paragraph._p.get_or_add_pPr()
            num = OxmlElement("w:numPr")
            num_id = OxmlElement("w:numId")
            num_id.set(qn("w:val"), "1")
            num.append(num_id)
            properties.append(num)
    parent.add_paragraph("Experience stays unchanged.")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def install(raw):
    source = config.seed_dir() / config.seed_files()["cv"]
    source.parent.mkdir(parents=True, exist_ok=True)
    source.write_bytes(raw)
    docs.ensure_templates()


@pytest.mark.parametrize("count", [3, 5, 8])
@pytest.mark.parametrize("style", ["List Bullet", "List Number", "Numbered", "Inherited"])
def test_supported_bullets_and_counts_render_exactly(count, style, tmp_path):
    raw = cv_bytes(count, style, table=True)
    inspection = docs.inspect_cv(raw)
    assert inspection["strong_count"] == count
    install(raw)
    assert docs.cv_identity()["strong_count"] == count
    tailoring = {
        "profile": "Public profile. " * 12,
        "strong_points": [f"Tailored {i}" for i in range(count)],
    }
    parsed = docs.parse_cv_tailoring(json.dumps(tailoring), count)
    output = tmp_path / "rendered.docx"
    docs.render_cv({}, parsed, output)
    text = "\n".join(p.text for p in docs._all_paragraphs(Document(output)))
    assert "{{" not in text
    assert all(point in text for point in tailoring["strong_points"])
    assert "Experience stays unchanged." in text


@pytest.mark.parametrize("count", [2, 9])
def test_invalid_count_rejected_without_writing(count):
    with pytest.raises(ValueError, match="3–8"):
        cv_uploads.stage(cv_bytes(count), "invalid.docx", "session")
    assert not config.user_path(docs.CV_TEMPLATE, "templates").exists()
    assert not (config.seed_dir() / config.seed_files()["cv"]).exists()


def test_existing_template_recovers_old_identity_count_and_rejects_mismatch(tmp_path):
    template, identity = docs.prepare_template(cv_bytes(4))
    docs.CV_TEMPLATE.parent.mkdir(parents=True)
    docs.CV_TEMPLATE.write_bytes(template)
    docs.IDENTITY_FILE.write_text(json.dumps({"name": identity["name"]}))
    docs.ensure_templates()
    assert docs.cv_identity()["strong_count"] == 4
    with pytest.raises(ValueError, match="exactly 4"):
        docs.render_cv(
            {}, {"profile": "x" * 200, "strong_points": ["a"] * 5}, tmp_path / "bad.docx"
        )
    document = Document(io.BytesIO(template))
    document.add_paragraph("{{STRONG_1}}")
    with pytest.raises(ValueError, match="unique contiguous"):
        docs.template_count(document)


def test_generation_sends_exact_schema_and_records_usage_before_bad_count(monkeypatch):
    install(cv_bytes(4))
    job = store.get_job(add_job())
    manager = llm.ServerManager()
    manager.client = llm.FakeLLMClient(
        [json.dumps({"profile": "x" * 200, "strong_points": ["a"] * 5})]
    )
    monkeypatch.setattr(docs, "convert_pdf", lambda path: None)
    with pytest.raises(ValueError, match="exactly 4"):
        docs.generate_for_job(manager, job, "resume")
    schema = manager.client.calls[0]["schema"]
    assert schema["properties"]["strong_points"]["minItems"] == 4
    assert schema["properties"]["strong_points"]["maxItems"] == 4
    assert schema["additionalProperties"] is False
    assert store.usage_total()["calls"] == 1
    assert store.list_documents(job["id"]) == []


def test_preview_confirmation_is_owner_session_bound_and_not_replayable():
    with tenancy.scope(A):
        preview = cv_uploads.stage(cv_bytes(3), "<unsafe>.docx", "session-a")
        assert not config.user_path(docs.CV_TEMPLATE, "templates").exists()
        with pytest.raises(ValueError, match="niet beschikbaar"):
            cv_uploads.confirm(preview["token"], "other-session")
    with tenancy.scope(B):
        with pytest.raises(ValueError, match="niet beschikbaar"):
            cv_uploads.confirm(preview["token"], "session-a")
    with tenancy.scope(A):
        assert cv_uploads.confirm(preview["token"], "session-a")["strong_count"] == 3
        assert docs.cv_identity()["strong_count"] == 3
        with pytest.raises(ValueError, match="niet beschikbaar"):
            cv_uploads.confirm(preview["token"], "session-a")


def test_preview_expiry_integrity_and_cancel(monkeypatch):
    monkeypatch.setattr(cv_uploads.time, "time", lambda: 1800000000)
    preview = cv_uploads.stage(cv_bytes(), "own.docx", "session")
    monkeypatch.setattr(cv_uploads.time, "time", lambda: 1800001000)
    with pytest.raises(ValueError, match="verlopen"):
        cv_uploads.confirm(preview["token"], "session")
    second = cv_uploads.stage(cv_bytes(), "own.docx", "session")
    cv_uploads.discard(second["token"], "session")
    with pytest.raises(ValueError):
        cv_uploads.confirm(second["token"], "session")
    third = cv_uploads.stage(cv_bytes(), "own.docx", "session")
    (cv_uploads._directory() / (third["token"] + ".docx")).write_bytes(b"tampered")
    with pytest.raises(ValueError):
        cv_uploads.confirm(third["token"], "session")


def test_concurrent_confirmations_have_one_winner_without_overwrite():
    with tenancy.scope(A):
        previews = [cv_uploads.stage(cv_bytes(count), "own.docx", "session") for count in (3, 8)]
    ready = threading.Barrier(2)

    def confirm(preview):
        with tenancy.scope(A):
            ready.wait(timeout=3)
            try:
                return cv_uploads.confirm(preview["token"], "session")["strong_count"]
            except FileExistsError:
                return None

    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(confirm, previews))
    assert results.count(None) == 1
    winner = next(item for item in results if item)
    with tenancy.scope(A):
        source = config.seed_dir() / config.seed_files()["cv"]
        assert docs.inspect_cv(source.read_bytes())["strong_count"] == winner
        assert docs.cv_identity()["strong_count"] == winner


def test_failed_publication_rolls_back_only_its_created_files(monkeypatch):
    preview = cv_uploads.stage(cv_bytes(), "own.docx", "session")
    original = cv_uploads._publish_exclusive
    calls = []

    def fail_second(path, raw):
        calls.append(path)
        if len(calls) == 2:
            raise OSError("synthetic disk failure")
        original(path, raw)

    monkeypatch.setattr(cv_uploads, "_publish_exclusive", fail_second)
    with pytest.raises(OSError):
        cv_uploads.confirm(preview["token"], "session")
    assert not (config.seed_dir() / config.seed_files()["cv"]).exists()
    assert not docs.CV_TEMPLATE.exists()
    monkeypatch.setattr(cv_uploads, "_publish_exclusive", original)
    cv_uploads.confirm(preview["token"], "session")


def test_dynamic_schema_is_locally_exact():
    with pytest.raises(ValueError):
        llm_schemas.cv_tailoring_model(2)
    with pytest.raises(ValueError):
        docs.parse_cv_tailoring(json.dumps({"profile": "x" * 200, "strong_points": ["a"] * 3}), 8)
