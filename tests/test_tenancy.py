"""Cross-account isolation and legacy preservation using synthetic data only."""

import io
import json
from concurrent.futures import ThreadPoolExecutor
from types import SimpleNamespace

import pytest
from bs4 import BeautifulSoup
from docx import Document
from fastapi.testclient import TestClient

from app import (
    auth,
    config,
    docs,
    feedback,
    geo,
    interview_state,
    llm,
    main,
    matching,
    profile,
    ranking,
    standards,
    store,
    systemone,
    tenancy,
)

A = {"id": "owner-a", "username": "a", "display_name": "Account A", "role": "user"}
B = {"id": "owner-b", "username": "b", "display_name": "Account B", "role": "user"}


@pytest.fixture()
def isolated(tmp_path, monkeypatch):
    for name, value in {
        "DATA_DIR": tmp_path,
        "DB_PATH": tmp_path / "app.db",
        "PROFILE_DIR": tmp_path / "profile",
        "JOBS_DIR": tmp_path / "jobs",
        "RUNTIME_DIR": tmp_path / "runtime",
        "SEED_DIR": tmp_path / "seed",
    }.items():
        monkeypatch.setattr(config, name, value)
    for module, names, section in [
        (
            profile,
            [
                "PROFILE_FILE",
                "PREFERENCES_FILE",
                "ANSWER_LOG_FILE",
                "INTERVIEW_PLAN_FILE",
                "HISTORY_DIR",
                "DRAFTS_DIR",
            ],
            "profile",
        ),
        (interview_state, ["STATE_FILE", "PLAN_FILE", "ANSWER_LOG_FILE"], "profile"),
        (feedback, ["FEEDBACK_FILE"], "profile"),
        (docs, ["CV_TEMPLATE", "IDENTITY_FILE"], "templates"),
    ]:
        for name in names:
            monkeypatch.setattr(module, name, tmp_path / section / getattr(module, name).name)
    monkeypatch.setattr(config, "APP_USER", "owner")
    monkeypatch.setattr(config, "APP_PASSWORD", "local-fixture-password")
    store.init_db()
    auth.init_auth()
    return tmp_path


def job():
    return store.upsert_job(
        {
            "url": "https://example.test/public",
            "title": "Public job",
            "description": "A public description",
            "company": "Public employer",
            "content_hash": "public",
        }
    )[0]


def test_catalog_shared_but_votes_matches_status_and_documents_are_private(isolated):
    jid = job()
    with tenancy.scope(A):
        store.set_job_vote(jid, 1)
        store.update_job_status(jid, "saved")
        store.update_job_match(jid, 88, "A_ONLY_REASON", "digest-a")
        store.update_job_signals(jid, 77, '{"fit":90}', "signals-a", 0.9)
        doc_id = store.add_document(jid, "resume", 1, "a.docx", "a.pdf", "{}")
        assert store.get_job(jid)["vote"] == 1
    with tenancy.scope(B):
        row = store.get_job(jid)
        assert row["description"] == "A public description"
        assert row["vote"] == 0 and row["status"] == "new"
        assert row["match_score"] is None and row["typed_signals"] is None
        assert store.get_document(doc_id) is None
        assert store.list_documents(jid) == []
        assert store.list_votes() == {"up": [], "down": []}
        assert store.count_documents() == 0
        store.add_document(jid, "resume", 1, "b.docx", "b.pdf", "{}")
    with tenancy.scope(A):
        assert store.get_job(jid)["match_reason"] == "A_ONLY_REASON"
        assert len(store.list_documents(jid)) == 1


def test_signal_cache_tracks_own_preferences_but_not_votes(isolated):
    jid = job()
    with tenancy.scope(A):
        row = store.get_job(jid)
        original = systemone.signals_key(row, "Own public profile")
        store.set_job_vote(jid, 1)
        assert systemone.signals_key(row, "Own public profile") == original
        profile.write_preferences({"hard_filters": {"no_phone_main": True}})
        assert systemone.signals_key(row, "Own public profile") != original
    with tenancy.scope(B):
        assert systemone.signals_key(row, "Own public profile") == original


def test_sessions_messages_usage_versions_and_meta_private(isolated):
    with tenancy.scope(A):
        store.create_session("session-a", "A private title", "interviewer", "fixture")
        store.add_message("session-a", "msg-a", "user", "A_SECRET")
        store.record_usage("session-a", "interview", 10, 5, 1)
        store.add_profile_version("a-history.md", "private")
        store.meta_set("consolidation_draft", "a-draft.json")
    with tenancy.scope(B):
        assert store.get_session("session-a") is None
        assert store.list_sessions() == [] and store.list_messages("session-a") == []
        assert store.usage_total()["calls"] == 0
        assert store.list_profile_versions() == []
        assert store.meta_get("consolidation_draft") is None
        store.update_session("session-a", title="SHOULD_NOT_CHANGE")
    with tenancy.scope(A):
        assert store.get_session("session-a")["title"] == "A private title"


def test_private_files_and_geo_cache_never_reuse_another_owner(isolated):
    with tenancy.scope(A):
        profile.write_profile("## Werk\nA_ONLY_PROFILE", "fixture")
        profile.write_preferences({"home": {"label": "A home", "lat": 1, "lon": 2}})
        profile.append_answer("Question", "A_ONLY_ANSWER")
        interview_state.ensure_state()
        assert geo.home()["label"] == "A home"
    with tenancy.scope(B):
        assert profile.read_profile() == ""
        assert profile.read_answer_log() == ""
        assert not profile.seed_available()
        assert "A_ONLY" not in str(profile.read_preferences())
        assert geo.home()["label"] == "Nog niet ingesteld"
        assert geo.assess("Amsterdam")["in_range"] is None
        profile.write_preferences({})
        prefs = profile.read_preferences()
        assert prefs["home"]["lat"] == 0
        assert not prefs["hard_filters"]["no_phone_main"]
        block = standards.standards_block(("cv-brief-standaarden",))
        assert "24-32" not in block
        assert "A_ONLY" not in profile.build_consolidation_prompt()


def test_background_workers_capture_owner_without_cross_account_leak(isolated):
    with tenancy.scope(A):
        worker_a = tenancy.capture(lambda: (tenancy.user_id(), store.meta_set("fixture", "A")))
    with tenancy.scope(B):
        worker_b = tenancy.capture(lambda: (tenancy.user_id(), store.meta_set("fixture", "B")))
    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(lambda f: f(), [worker_a, worker_b]))
    assert [r[0] for r in results] == ["owner-a", "owner-b"]
    with tenancy.scope(A):
        assert store.meta_get("fixture") == "A"
    with tenancy.scope(B):
        assert store.meta_get("fixture") == "B"
    assert tenancy.current() is None


def test_migration_is_idempotent_and_keeps_legacy_state(isolated):
    jid = job()
    store.set_job_vote(jid, -1)
    store.update_job_match(jid, 42, "LEGACY_REASON", "original_digest")
    store.meta_set("profile_bootstrapped_at", "partially_bootstrapped")
    profile.write_profile("## Werk\nORIGINAL_LEGACY_PROFILE", "fixture")
    before = profile.PROFILE_FILE.read_bytes()
    store.init_db()
    store.init_db()
    assert profile.PROFILE_FILE.read_bytes() == before
    assert store.get_job(jid)["vote"] == -1
    assert store.get_job(jid)["profile_hash"] == "original_digest"
    assert store.meta_get("profile_bootstrapped_at") == "partially_bootstrapped"
    with tenancy.scope(A):
        assert store.get_job(jid)["match_score"] is None
        assert store.meta_get("profile_bootstrapped_at") is None
        assert profile.read_profile() == ""


def test_recovery_repairs_all_owners_not_just_legacy(isolated):
    for owner in [A, B]:
        with tenancy.scope(owner):
            store.meta_set("match_status", "running")
            store.meta_set("doc_status:1", '{"state":"running"}')
    store.recover_interrupted_work()
    for owner in [A, B]:
        with tenancy.scope(owner):
            assert store.meta_get("match_status") == "interrupted"
            assert json.loads(store.meta_get("doc_status:1"))["state"] == "error"


class FakeGenerator:
    def __init__(self, text="not valid json", usage=None):
        self.text = text
        self.usage = usage or {"input_tokens": 120, "output_tokens": 8, "cost_usd": 0.01}

    def generate(self, *args, **kwargs):
        return self.text, self.usage


def test_failed_verdict_still_records_paid_provider_usage(isolated, monkeypatch):
    with tenancy.scope(A):
        jid = job()
        profile.write_profile("## Werk\nPublic profile", "fixture")
        monkeypatch.setattr(systemone, "enabled", lambda: False)
        monkeypatch.setattr(
            feedback,
            "load",
            lambda: {"counts": {"up": 0, "down": 0}, "liked_keywords": [], "disliked_keywords": []},
        )
        monkeypatch.setattr(ranking, "rerank_all", lambda *args, **kwargs: 0)
        monkeypatch.setattr(
            ranking, "escalation_candidates", lambda *args, **kwargs: [store.get_job(jid)]
        )
        monkeypatch.setattr(ranking, "queue_depth", lambda *args: 0)
        summary = matching.match_jobs(FakeGenerator(), limit=1)
        assert summary["failed"] == 1
        assert summary["cost_usd"] == 0.01
        assert store.usage_total()["cost"] == 0.01
        assert store.usage_total()["calls"] == 1


def test_failed_document_still_records_usage_and_never_overwrites_orphan(isolated, monkeypatch):
    with tenancy.scope(A):
        jid = job()
        profile.write_profile("## Werk\nPublic profile", "fixture")
        directory = docs._job_dir(store.get_job(jid))
        orphan = directory / "resume-v1.docx"
        orphan.write_bytes(b"ORIGINAL_UNREGISTERED_DOCUMENT")
        monkeypatch.setattr(docs, "ensure_templates", lambda: None)
        monkeypatch.setattr(docs, "cv_prompt", lambda *args: "prompt")
        monkeypatch.setattr(docs, "parse_cv_tailoring", lambda text: {})
        monkeypatch.setattr(docs, "render_cv", lambda job, data, path: path.write_bytes(b"NEW"))
        monkeypatch.setattr(docs, "convert_pdf", lambda path: None)
        result = docs.generate_for_job(FakeGenerator('{"ok": true}'), store.get_job(jid), "resume")
        assert result["version"] == 2
        assert orphan.read_bytes() == b"ORIGINAL_UNREGISTERED_DOCUMENT"
        assert store.usage_total()["cost"] == 0.01

        def fail_parse(text):
            raise ValueError("malformed provider answer")

        monkeypatch.setattr(docs, "parse_cv_tailoring", fail_parse)
        with pytest.raises(ValueError):
            docs.generate_for_job(FakeGenerator("bad"), store.get_job(jid), "resume")
        assert store.usage_total()["cost"] == 0.02
        assert store.usage_total()["calls"] == 2


def test_quality_helper_reads_owner_scoped_scores(isolated):
    with tenancy.scope(A):
        jid = job()
        store.update_job_match(jid, 80, "A only", "digest")
        store.set_rank_scores({jid: 77})
        assert [row["rank_score"] for row in store.list_scored_jobs()] == [77]
    with tenancy.scope(B):
        assert store.list_scored_jobs() == []


def test_interview_prompt_only_contains_owners_public_profile(isolated):
    with tenancy.scope(A):
        profile.write_profile("## Werk\nOTHER_OWNER_SECRET", "fixture")
    with tenancy.scope(B):
        profile.write_profile(
            "## Werk\nPUBLIC_B\n## Privé\nPRIVATE_B_SENTINEL\n## Open vragen\nOPEN_B_SENTINEL",
            "fixture",
        )
        profile.append_answer("Private", "ANSWER_LOG_SENTINEL")
        sent = []
        manager = llm.ServerManager()
        manager.client = SimpleNamespace(
            session_info=lambda _: {},
            prompt=lambda _, text: sent.append(text),
            wait=lambda _: None,
            messages=lambda _: [
                {
                    "type": "assistant",
                    "id": "msg",
                    "content": [{"type": "text", "text": "Wat zoek je?"}],
                }
            ],
        )
        manager.run_turn("fixture-session", "Hallo")
        assert "PUBLIC_B" in sent[0]
        assert all(
            secret not in sent[0]
            for secret in [
                "OTHER_OWNER_SECRET",
                "PRIVATE_B_SENTINEL",
                "OPEN_B_SENTINEL",
                "ANSWER_LOG_SENTINEL",
            ]
        )


def login(client, username, password):
    token = BeautifulSoup(client.get("/login").text, "html.parser").select_one(
        'input[name="_csrf"]'
    )["value"]
    result = client.post(
        "/login",
        data={"username": username, "password": password, "_csrf": token},
        follow_redirects=False,
    )
    assert result.status_code == 303
    html = BeautifulSoup(client.get("/account").text, "html.parser")
    client.headers["x-csrf-token"] = html.select_one('meta[name="csrf-token"]')["content"]


def test_own_cv_upload_cannot_read_or_replace_owners_template(isolated):
    credentials = json.loads((config.RUNTIME_DIR / "bootstrap-accounts.json").read_text())[
        "accounts"
    ]
    client = TestClient(main.app)
    login(client, "test-user", credentials["test-user"]["password"])
    token = client.headers.pop("x-csrf-token")
    document = Document()
    document.add_paragraph("Synthetic Test Person")
    document.add_paragraph("Own career direction")
    document.add_paragraph("test@example.test")
    document.add_paragraph("PROFIEL")
    document.add_paragraph("Original public experience from this user.")
    document.add_paragraph("STERKE PUNTEN")
    for i in range(5):
        document.add_paragraph(f"Own skill {i}", style="List Bullet")
    content = io.BytesIO()
    document.save(content)
    response = client.post(
        "/account/cv",
        data={"_csrf": token},
        files={
            "cv": (
                "own.docx",
                content.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        follow_redirects=False,
    )
    assert response.status_code == 303
    uid = auth._get_user_by_username("test-user")["id"]
    with tenancy.scope({"id": uid, "username": "test-user"}):
        assert (config.seed_dir() / "cv.docx").exists()
        assert "Synthetic Test Person" in docs.cv_identity()["name"]
    assert not docs.CV_TEMPLATE.exists()
    assert not (config.SEED_DIR / config.SEED_FILES["cv"]).exists()
    client.close()


def test_routes_deny_cross_account_ids_and_admin_privileges(isolated):
    with tenancy.scope({"id": "legacy"}):
        jid = job()
        store.create_session("private-owner-session", "PRIVATE_TITLE", "interviewer", "fixture")
        store.add_message("private-owner-session", "message", "user", "PRIVATE_MESSAGE")
        doc_id = store.add_document(jid, "resume", 1, "jobs/private.docx", "jobs/private.pdf", "{}")
        profile.write_profile("## Werk\nLEGACY_PRIVATE_PROFILE", "fixture")
    credentials = json.loads((config.RUNTIME_DIR / "bootstrap-accounts.json").read_text())[
        "accounts"
    ]
    client = TestClient(main.app)
    login(client, "test-user", credentials["test-user"]["password"])
    for url in [
        f"/jobs/{jid}/documents/{doc_id}/pdf",
        "/api/sessions/private-owner-session/messages",
        "/interview?session=private-owner-session",
    ]:
        assert client.get(url).status_code == 404
    assert (
        client.post(
            "/api/sessions/private-owner-session/send", json={"text": "intrusion"}
        ).status_code
        == 404
    )
    assert client.get("/users").status_code == 403 and client.get("/admin").status_code == 403
    assert client.post("/sources/scan").status_code == 403
    assert "LEGACY_PRIVATE_PROFILE" not in client.get("/profile").text
    assert client.get("/jobs").status_code == 200
    assert client.get("/api/sessions").json() == {"sessions": []}
    assert client.post("/profile/restore/1").status_code == 404
    assert (
        client.post(
            "/jobs/1/vote", data={"vote": "up"}, headers={"x-csrf-token": "incorrect"}
        ).status_code
        == 403
    )
    admin = TestClient(main.app)
    login(admin, "admin", credentials["admin"]["password"])
    assert admin.get("/users").status_code == 200
    assert "LEGACY_PRIVATE_PROFILE" not in admin.get("/admin").text
    client.close()
    admin.close()
